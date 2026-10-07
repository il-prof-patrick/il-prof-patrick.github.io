"""Correzioni al Word di brainstorming (6 ottobre 2026): Ricorda riscritti e animazioni
allineate al testo. Uso: python3 correggi_word.py <entrata.docx> <uscita.docx>"""
import sys, re, copy, docx
from docx.oxml.ns import qn

d = docx.Document(sys.argv[1])
T = d.tables

def sostituisci(sub, nuovo, cerca_in=None):
    """Sostituisce `sub` dentro un singolo run (conserva la formattazione)."""
    trovati = 0
    tabs = [T[i] for i in cerca_in] if cerca_in else T
    for t in tabs:
        for p in (q for r in t.rows for q in r.cells[0].paragraphs):
            for r in p.runs:
                if sub in r.text:
                    r.text = r.text.replace(sub, nuovo); trovati += 1
    assert trovati == 1, (sub, trovati)

# ---------------------------------------------------------------- animazioni e schemi
sostituisci(', GAN (il globo)', '')                                                        # F2.2.1
sostituisci(' (nelle versioni a doppio anello il traffico torna indietro e la rete resiste)', '')  # F2.3.2
sostituisci(' (il gateway)', '')                                                           # F2.6.1
sostituisci('grida in broadcast', 'chiede a tutta la rete')                                # F3.5.1
sostituisci('la centrale dell’ISP, il MIX di Milano,', 'la centrale dell’ISP, Milano,')     # F4.2.8
sostituisci('fino al MIX di Milano,', 'fino a Milano,')
sostituisci(': si scambiano le chiavi e compare', ': browser e server si accordano per cifrare la comunicazione (HTTPS) e compare')  # F5.5.1
sostituisci('da sensore ad attuatore passando per la rete', 'dal sensore che misura alla caldaia che si accende, passando per la rete')  # F9.6.2
sostituisci('microcontrollore confronta con l’obiettivo (20 °C) e, tramite il relè (attuatore), accende', 'termostato confronta il valore con l’obiettivo (20 °C) e accende')
sostituisci('una parola (un token) alla volta', 'una parola alla volta')                   # F10.1.2

# F10.5.1 (apprendimento federato): il testo non lo spiega più → via tutta l'animazione
fed = [t for t in T if 'F10.5.1' in t.cell(0, 0).paragraphs[0].text]
assert len(fed) == 1
el = fed[0]._tbl; prev = el.getprevious()
el.getparent().remove(el)
if prev is not None and prev.tag == qn('w:p') and not ''.join(prev.itertext()).strip():
    prev.getparent().remove(prev)

# ---------------------------------------------------------------- Ricorda
RICORDA = {
 2: [
  0,
  'Per estensione le reti si dividono in **BAN** (corpo), **PAN** (persona), **LAN** (edificio), **MAN** (città) e **WAN** (regioni, nazioni, continenti); conta l’**estensione geografica**, non il numero di dispositivi.',
  2, 3,
  'La **scheda di rete** ha un **indirizzo MAC** unico; l’**hub** ripete i dati a tutte le porte, lo **switch** li consegna solo al destinatario; il **router** collega reti diverse e sceglie la strada usando gli **indirizzi IP**; il **modem** trasforma i bit in un segnale adatto alla linea dell’operatore e viceversa.',
  'L’**access point** crea il Wi-Fi; i **ripetitori** lo estendono ma spesso dimezzano la banda, i sistemi **mesh** creano un’unica rete in tutta la casa. Lo **switch** lavora **dentro una rete**, il **router tra reti**; il “router” di casa è un **tutto-in-uno**: modem, router, switch e access point.',
 ],
 3: [
  'Un **protocollo** è l’insieme di regole con cui due dispositivi si scambiano i dati: **sintassi** (il formato), **semantica** (il significato) e **sequenza** (chi parla e quando).',
  'La comunicazione è divisa in **livelli**: ognuno risolve un problema e si può cambiare senza toccare gli altri. Internet usa il modello **TCP/IP**, con **quattro** livelli: accesso alla rete, Internet, trasporto, applicazione.',
  2, 3,
  'Il **DHCP** assegna automaticamente a chi si collega un **indirizzo IP** libero, insieme all’indirizzo del router (*gateway*) e del server DNS.',
  'Il **DNS** traduce i **nomi** dei siti in **indirizzi IP**; è **distribuito e gerarchico** (i nomi si leggono da destra) e usa la **cache** per rispondere subito alle richieste già fatte.',
  5,
 ],
 4: [
  '**Internet** è una **rete di reti**: decine di migliaia di reti indipendenti che comunicano con i protocolli **TCP/IP**. È nata da **ARPANET**, il cui primo messaggio partì nel **1969**.',
  'Fisicamente Internet è fatta di **ISP** (fornitori di accesso), **dorsali** in fibra, **cavi sottomarini** (oltre il 95% del traffico tra continenti) e **data center**.',
  2, 3, 4,
  '**Internet** è l’infrastruttura (rete di reti, TCP/IP); il **Web** è **uno dei servizi** che funzionano sopra Internet: pagine collegate da link, consultate con un browser. Altri servizi (posta, chat, giochi) usano Internet senza il Web.',
  'Il Web nasce al **CERN** di Ginevra da **Tim Berners-Lee** (proposta nel 1989, prima pagina nel 1991); nel 1993 il CERN lo rende **pubblico e gratuito**, e per questo si diffonde in fretta.',
  6,
  'L’**URL** è l’indirizzo di una risorsa: **protocollo, dominio, percorso, parametri, segnalibro**. Il dominio si legge **da destra**: conta il **dominio registrato** (es. poste.it), non le parole che compaiono altrove nell’indirizzo.',
  '**HTTP** funziona a **richiesta e risposta** con **codici di stato** (2xx successo, 3xx reindirizzamento, 4xx errore del client, es. 404; 5xx errore del server). **HTTPS** cifra la comunicazione con TLS; il **lucchetto** garantisce la cifratura, **non l’onestà del sito**.',
  9, 10, 11, 12, 13,
 ],
 5: [
  0,
  'L’ordine dei risultati dipende da pertinenza, autorevolezza, data, qualità tecnica e **da chi cerca** (lingua, posizione, ricerche precedenti): i primi sono i più pertinenti e popolari, e **popolarità non vuol dire verità**.',
  'I risultati **organici** sono scelti dall’algoritmo, quelli **sponsorizzati** sono **annunci pagati**: la pubblicità è ciò che rende gratuita la ricerca.',
  'Per **cercare bene** si usano parole chiave specifiche e i filtri del motore; la **ricerca inversa** delle immagini (ad esempio con Google Lens) mostra dove e da quando compare una foto.',
  3,
  'La **misinformazione** è falsa ma condivisa in buona fede, la **disinformazione** è creata apposta; fake news, clickbait, contenuti fuori contesto e falsi creati con l’IA si diffondono sfruttando le **emozioni**. Il **fact-checking** verifica risalendo alle fonti originali.',
  'Il **feed** è scelto da un **algoritmo di raccomandazione** che premia ciò che ci trattiene (*engagement*); così si finisce nella **bolla dei filtri**, che mostra soprattutto ciò che conferma le nostre idee.',
  5,
 ],
 6: [
  0,
  'Le **generazioni**: **1G** analogica, solo voce; **2G** digitale e SMS; **3G** Internet sul telefono; **4G** Internet mobile veloce (streaming, social, app); **5G** dal **2019**, per latenza bassissima e moltissimi dispositivi.',
  2,
  'L’**Internet of Things** è la rete di **oggetti** dotati di **sensori**, capacità di calcolo e connessione, che raccolgono dati, li scambiano via Internet e possono **agire da soli**.',
  'L’IoT è già ovunque: nella **casa intelligente** (risparmio energetico, autonomia di anziani e persone con disabilità) e nella **città intelligente** (parcheggi, lampioni, semafori, qualità dell’aria).',
  'Rischi dell’IoT: **sicurezza** (password predefinite, aggiornamenti che non arrivano, dati non cifrati: nel **2016** la botnet **Mirai** rese irraggiungibili molti siti famosi) e **privacy** (dati molto personali, riprese di altre persone).',
  5, 6, 7,
 ],
 7: [
  0,
  'I modelli linguistici imparano da migliaia di miliardi di parole raccolte dal **Web** con i **crawler**. L’uso dei testi solleva problemi di **diritto d’autore**: in Europa i titolari dei diritti possono **opporsi** all’estrazione.',
  'L’**IA sul dispositivo** (*edge AI*, con le **NPU**) è più veloce, più riservata e funziona senza Internet, ma è meno potente di quella nel **cloud**; spesso le due si combinano.',
  'La **teoria dell’Internet morta** sostiene che gran parte di ciò che si trova online sia prodotta da **bot** e programmi automatici, e che le persone vere siano ormai una minoranza.',
  'Anche senza complotto, le conseguenze sono reali: numeri e recensioni non sono più prove, l’opinione pubblica può essere falsata, la qualità del Web cala e i modelli addestrati su testi generati rischiano il **collasso del modello**. La risposta è **verificare** e **dare valore ai contenuti umani** e alle fonti riconoscibili.',
 ],
}
# Nella lista un numero k = punto k originale del Word, invariato.

ric = [r.cells[0] for t in T for r in t.rows if r.cells[0].paragraphs[0].text.strip().startswith('RICORDA')]
assert len(ric) == 7, len(ric)

def riempi(p, testo, modello):
    """Svuota il paragrafo e lo riscrive: **grassetto**, *corsivo*."""
    for r in list(p.runs): r._r.getparent().remove(r._r)
    base = None
    for r in modello.runs:
        if not r.bold: base = r; break
    for pezzo in re.split(r'(\*\*.+?\*\*|\*[^*]+?\*)', testo):
        if not pezzo: continue
        b = pezzo.startswith('**'); i = not b and pezzo.startswith('*')
        txt = pezzo.strip('*') if (b or i) else pezzo
        r = p.add_run(txt)
        if base is not None and base._r.rPr is not None:
            r._r.insert(0, copy.deepcopy(base._r.rPr))
        r.bold = True if b else None
        r.italic = True if i else None

for n, nuovi in RICORDA.items():
    cella = ric[n - 1]
    vecchi = [p for p in cella.paragraphs[1:] if p.text.strip()]
    modello = vecchi[0]
    risultato = []
    for x in nuovi:
        if isinstance(x, int): risultato.append(vecchi[x])
        else:
            np_ = copy.deepcopy(modello._p); nuovo = docx.text.paragraph.Paragraph(np_, cella)
            riempi(nuovo, x, modello); risultato.append(nuovo)
    # i vecchi non riusati spariscono; i nuovi vanno al posto della lista
    prec = cella.paragraphs[0]._p   # l'etichetta RICORDA
    for p in risultato:
        prec.addnext(p._p); prec = p._p
    for p in vecchi:
        if p._p.getparent() is not None and not any(p._p is q._p for q in risultato):
            p._p.getparent().remove(p._p)

d.save(sys.argv[2])
print('ok')
